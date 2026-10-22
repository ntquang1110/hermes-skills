#!/usr/bin/env python3
"""Sheets export fallback: dump every tab + formulas without Sheets API.

Usage:
    python sheets_export_fallback.py <SPREADSHEET_ID> [--out DIR]

Does in one run what the references/sheets-export-fallback.md recipe does:
  1. Fetch htmlview, enumerate all gids
  2. Export each gid to CSV (display values)
  3. Export full workbook to xlsx
  4. Dump every non-empty row of every tab (formulas preserved)

Requires: requests, openpyxl (pip install requests openpyxl).
See references/sheets-export-fallback.md for pitfalls (Windows /tmp, empty tabs).
"""
import re
import sys
import zipfile
import argparse
import xml.etree.ElementTree as ET
from pathlib import Path

try:
    import requests
except ImportError:
    sys.exit("Need 'requests': pip install requests openpyxl")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("sheet_id")
    ap.add_argument("--out", default="sheet_dump")
    args = ap.parse_args()

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    base = f"https://docs.google.com/spreadsheets/d/{args.sheet_id}"

    html = requests.get(f"{base}/htmlview", timeout=20).text
    gids = sorted({g for g in re.findall(r"gid=(\d+)", html)})
    print("GIDS:", gids)

    for gid in gids:
        r = requests.get(f"{base}/export?format=csv&gid={gid}", timeout=20)
        p = out / f"tab_{gid}.csv"
        p.write_bytes(r.content)
        print(f"gid={gid} csv_bytes={len(r.content)} -> {p}"
              + (" (EMPTY TAB)" if len(r.content) == 0 else ""))

    xlsx = out / "wb.xlsx"
    r = requests.get(f"{base}/export?format=xlsx", timeout=30)
    xlsx.write_bytes(r.content)
    print(f"xlsx_bytes={len(r.content)} -> {xlsx}")

    # Tab names from workbook.xml (no openpyxl needed for this part)
    with zipfile.ZipFile(xlsx) as z:
        root = ET.fromstring(z.read("xl/workbook.xml"))
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        print("TAB NAMES:", [s.get("name") for s in root.iter(ns + "sheet")])

    # Full content dump with formulas preserved
    try:
        import openpyxl
    except ImportError:
        print("openpyxl missing: skipping content dump (CSV/xlsx already saved).")
        print("pip install openpyxl to enable it.")
        return
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    print("SHEETS:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"===== {name}  rows={ws.max_row} cols={ws.max_column} =====")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() != "" for v in row):
                print(list(row))


if __name__ == "__main__":
    main()
